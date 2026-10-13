# -*- coding: utf-8 -*-
"""把 thesis.md 排版成符合中文本科毕业论文习惯的 .docx。

用法:
    python3 build_docx.py thesis.md 输出.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SONG = '宋体'
HEI = '黑体'
KAI = '楷体'
MONO = 'Consolas'

BODY_SIZE = Pt(12)        # 小四
CAPTION_SIZE = Pt(10.5)   # 五号
CODE_SIZE = Pt(9)
TABLE_SIZE = Pt(10.5)


def set_run_font(run, name=SONG, size=BODY_SIZE, bold=False, italic=False,
                 color=None):
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), name)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)


def add_rich(paragraph, text, name=SONG, size=BODY_SIZE, base_bold=False):
    """按 ** 分段写入，** 内加粗。"""
    for i, chunk in enumerate(re.split(r'\*\*', text)):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        set_run_font(run, name=name, size=size, bold=base_bold or (i % 2 == 1))


def style_paragraph(p, first_indent=True, space_after=6, align=None,
                    line_spacing=1.5):
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if first_indent:
        pf.first_line_indent = Pt(24)
    if align is not None:
        pf.alignment = align


def add_body(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, first_indent=True)
    add_rich(p, text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, first_indent=False, space_after=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
    add_rich(p, text, size=CAPTION_SIZE)
    return p


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        style_paragraph(p, first_indent=False, space_after=0,
                        line_spacing=1.0)
        pf = p.paragraph_format
        pf.left_indent = Pt(18)
        run = p.add_run(line if line else ' ')
        set_run_font(run, name=MONO, size=CODE_SIZE)
        # 浅灰底纹
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), 'F5F5F5')
        p._p.get_or_add_pPr().append(shd)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, first_indent=False, space_after=6, line_spacing=1.0)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        style_paragraph(p, first_indent=False, space_after=2, line_spacing=1.0,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(cell_text)
        set_run_font(run, name=HEI, size=TABLE_SIZE, bold=True)
    for row in body:
        cells = table.add_row().cells
        for i in range(len(header)):
            text = row[i] if i < len(row) else ''
            p = cells[i].paragraphs[0]
            style_paragraph(p, first_indent=False, space_after=2,
                            line_spacing=1.0)
            run = p.add_run(text)
            set_run_font(run, name=SONG, size=TABLE_SIZE)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, first_indent=False, space_after=8, line_spacing=1.0)


def add_image(doc, path, width_cm=14.5):
    if not path.is_file():
        add_body(doc, f'（缺少图片：{path}）')
        return
    p = doc.add_paragraph()
    style_paragraph(p, first_indent=False, space_after=4,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def add_heading(doc, text, level):
    style = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}[level]
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    if level == 1:
        pf.space_before = Pt(0)
        pf.space_after = Pt(18)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = Pt(16)
    elif level == 2:
        pf.space_before = Pt(12)
        pf.space_after = Pt(8)
        size = Pt(14)
    else:
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        size = Pt(12)
    run = p.add_run(text)
    set_run_font(run, name=HEI, size=size, bold=True,
                 color=RGBColor(0, 0, 0))
    return p


def add_page_number_footer(section, restart_at=1):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, name=SONG, size=Pt(10.5))
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn('w:pgNumType'))
    if pg_num_type is None:
        pg_num_type = OxmlElement('w:pgNumType')
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn('w:start'), str(restart_at))


def add_toc_field(doc, entries):
    """插入可更新的目录域；同时写入静态条目，未更新域时也能看到章节名。"""
    p = doc.add_paragraph()
    style_paragraph(p, first_indent=False, space_after=6, line_spacing=1.5)
    run = p.add_run()
    set_run_font(run, name=SONG, size=BODY_SIZE)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)

    p2 = doc.add_paragraph()
    style_paragraph(p2, first_indent=False, space_after=0, line_spacing=1.5)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run2 = p2.add_run()
    set_run_font(run2, name=SONG, size=BODY_SIZE)
    run2._r.append(end)

    # 静态条目插到 separate 与 end 之间
    anchor = p2._p
    for text, level in entries:
        entry = doc.add_paragraph()
        style_paragraph(entry, first_indent=False, space_after=2,
                        line_spacing=1.5)
        entry.paragraph_format.left_indent = Pt(14 * (level - 1))
        run = entry.add_run(text)
        set_run_font(run, name=SONG, size=Pt(12) if level == 1 else Pt(11),
                     bold=(level == 1))
        anchor.addprevious(entry._p)


def parse_markdown(text):
    """把 markdown 拆成块。"""
    blocks = []
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', code))
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-{2,}', c or '--') for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
            continue
        if line.startswith('!['):
            m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            blocks.append(('image', (m.group(1), m.group(2))))
            i += 1
            continue
        if line.startswith('#### '):
            blocks.append(('h3', line[5:].strip()))
            i += 1
            continue
        if line.startswith('### '):
            blocks.append(('h2', line[4:].strip()))
            i += 1
            continue
        if line.startswith('## '):
            blocks.append(('h1', line[3:].strip()))
            i += 1
            continue
        if line.startswith('# '):
            blocks.append(('title', line[2:].strip()))
            i += 1
            continue
        if line.startswith('> '):
            quote = []
            while i < len(lines) and (lines[i].startswith('> ') or lines[i] == '>'):
                quote.append(lines[i][2:] if lines[i].startswith('> ') else '')
                i += 1
            blocks.append(('quote', quote))
            continue
        if line.strip() == '---':
            i += 1
            continue
        if line.strip() == '':
            i += 1
            continue
        blocks.append(('p', line.strip()))
        i += 1
    return blocks


def build(md_path, out_path):
    text = Path(md_path).read_text(encoding='utf-8')
    blocks = parse_markdown(text)

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = SONG
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    # ---- 封面 ----
    title = ''
    info_lines = []
    for kind, value in blocks:
        if kind == 'title':
            title = value
        elif kind == 'p' and value.startswith('**') and '：' in value:
            info_lines.append(value)
        elif kind == 'h1':
            break

    def center_line(text, name, size, bold=False, space_after=0, space_before=0):
        p = doc.add_paragraph()
        style_paragraph(p, first_indent=False, space_after=space_after,
                        align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5)
        p.paragraph_format.space_before = Pt(space_before)
        add_rich(p, text, name=name, size=size, base_bold=bold)
        return p

    center_line('', SONG, BODY_SIZE, space_after=0, space_before=60)
    center_line('上海第二工业大学', HEI, Pt(26), bold=True, space_after=18)
    center_line('本科毕业论文（设计）', HEI, Pt(22), bold=True, space_after=48)
    center_line('题目：' + title, HEI, Pt(15), bold=True, space_after=60)
    for line in info_lines:
        center_line(line.replace('**', ''), SONG, Pt(14), space_after=6)
    center_line('', SONG, BODY_SIZE, space_after=0, space_before=36)

    doc.add_section(WD_SECTION.NEW_PAGE)
    body_section = doc.sections[-1]
    body_section.page_width = Cm(21)
    body_section.page_height = Cm(29.7)
    body_section.top_margin = Cm(2.54)
    body_section.bottom_margin = Cm(2.54)
    body_section.left_margin = Cm(3.0)
    body_section.right_margin = Cm(2.6)
    add_page_number_footer(body_section, restart_at=1)
    doc.sections[0].footer.is_linked_to_previous = False

    # 目录条目（来自 H1 / H2）
    toc_entries = []
    for kind, value in blocks:
        if kind == 'h1':
            if value == '目录':
                continue
            toc_entries.append((value, 1))
        elif kind == 'h2':
            toc_entries.append((value, 2))

    first_h1 = True
    skip_bullets = False
    base_dir = Path(md_path).parent
    for kind, value in blocks:
        if kind in ('title',):
            continue
        if kind == 'h1':
            if value == '目录':
                if not first_h1:
                    doc.add_page_break()
                first_h1 = False
                add_heading(doc, value, 1)
                add_toc_field(doc, toc_entries)
                skip_bullets = True
                continue
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            skip_bullets = False
            add_heading(doc, value, 1)
            continue
        if skip_bullets:
            continue
        if kind == 'h2':
            add_heading(doc, value, 2)
        elif kind == 'h3':
            add_heading(doc, value, 3)
        elif kind == 'p':
            if value.startswith('- '):
                p = doc.add_paragraph()
                style_paragraph(p, first_indent=False, space_after=4)
                p.paragraph_format.left_indent = Pt(24)
                add_rich(p, value[2:])
            elif value.startswith('**') and '：' in value:
                continue
            elif re.match(r'^(表|图)\s?\d', value):
                add_caption(doc, value)
            else:
                add_body(doc, value)
        elif kind == 'quote':
            for line in value:
                if not line.strip():
                    continue
                p = doc.add_paragraph()
                style_paragraph(p, first_indent=True, space_after=4,
                                line_spacing=1.5)
                p.paragraph_format.left_indent = Pt(24)
                add_rich(p, line, name=KAI)
        elif kind == 'code':
            add_code_block(doc, value)
        elif kind == 'table':
            if value:
                add_table(doc, value)
        elif kind == 'image':
            alt, rel = value
            add_image(doc, (base_dir / rel).resolve())

    doc.save(str(out_path))
    print(f'已生成 {out_path}')


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
